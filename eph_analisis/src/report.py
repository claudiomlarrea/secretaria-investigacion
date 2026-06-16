"""Exportación de resultados a Excel y Word."""

from __future__ import annotations

import io
import json
from pathlib import Path
from typing import Any

import pandas as pd

try:
    from docx import Document
    from docx.shared import Inches
except ImportError:  # pragma: no cover
    Document = None


def _escribir_excel(resultado: dict, destino) -> None:
    tablas: dict[str, pd.DataFrame] = resultado.get("tablas", {})
    with pd.ExcelWriter(destino, engine="openpyxl") as writer:
        meta = resultado.get("meta", {})
        pd.DataFrame([meta]).to_excel(writer, sheet_name="metadatos", index=False)
        for nombre, hoja in [
            ("descriptivos", "descriptivos_anuales"),
            ("frecuencias", "frecuencias"),
            ("correlaciones", "correlaciones"),
            ("logistica_coef", "logistica_coeficientes"),
            ("logistica_or", "logistica_odds_ratios"),
            ("shap", "shap_importancia"),
            ("cluster_tamanos", "cluster_tamanos"),
            ("cluster_perfiles", "cluster_perfiles"),
        ]:
            df = tablas.get(hoja)
            if df is not None and not df.empty:
                df.to_excel(writer, sheet_name=nombre[:31], index=False)
        modelos = resultado.get("modelos", {})
        if modelos:
            pd.DataFrame([modelos]).to_excel(writer, sheet_name="resumen_modelos", index=False)


def exportar_excel_bytes(resultado: dict) -> bytes:
    buf = io.BytesIO()
    _escribir_excel(resultado, buf)
    buf.seek(0)
    return buf.getvalue()


def exportar_excel(resultado: dict, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    _escribir_excel(resultado, path)
    return path


def _agregar_tabla(doc: Any, titulo: str, df: pd.DataFrame, max_filas: int = 30) -> None:
    doc.add_heading(titulo, level=2)
    if df is None or df.empty:
        doc.add_paragraph("Sin datos disponibles.")
        return

    sub = df.head(max_filas)
    table = doc.add_table(rows=1, cols=len(sub.columns))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, col in enumerate(sub.columns):
        hdr[i].text = str(col)

    for _, row in sub.iterrows():
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = "" if pd.isna(val) else str(round(val, 4) if isinstance(val, float) else val)


def exportar_word(
    resultado: dict,
    path: Path,
    *,
    titulo: str,
    periodo: str,
    ambito: str,
) -> Path:
    """Genera informe Word con resumen ejecutivo y tablas principales."""
    if Document is None:
        raise ImportError("Instale python-docx: pip install python-docx")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()

    doc.add_heading(titulo, level=0)
    doc.add_paragraph(f"Ámbito: {ambito}")
    doc.add_paragraph(f"Período: {periodo}")
    doc.add_paragraph(f"Fuente: INDEC — EPH (módulos hogar, individuo y TIC)")

    meta = resultado.get("meta", {})
    doc.add_heading("Resumen ejecutivo", level=1)
    doc.add_paragraph(
        f"Se analizaron {meta.get('registros', '—'):,} registros individuales (15+ años) "
        f"con variables proxy de exclusión digital y movilidad social."
    )

    corr = resultado.get("correlacion_destacada")
    if corr is not None:
        doc.add_paragraph(
            f"Correlación Pearson entre exclusión digital y score de movilidad proxy: {corr:.3f}"
        )

    tablas = resultado.get("tablas", {})
    _agregar_tabla(doc, "Indicadores anuales", tablas.get("descriptivos_anuales"))
    _agregar_tabla(doc, "Frecuencias ponderadas", tablas.get("frecuencias"))
    _agregar_tabla(doc, "Correlaciones", tablas.get("correlaciones"))

    modelos = resultado.get("modelos", {})
    if modelos:
        doc.add_heading("Modelos estadísticos", level=1)
        for clave, valor in modelos.items():
            if isinstance(valor, (dict, list)):
                texto = json.dumps(valor, ensure_ascii=False, indent=2)
            else:
                texto = str(valor)
            doc.add_paragraph(f"{clave}: {texto[:2000]}")

    _agregar_tabla(doc, "Regresión logística — coeficientes", tablas.get("logistica_coeficientes"))
    _agregar_tabla(doc, "Regresión logística — odds ratios", tablas.get("logistica_odds_ratios"))
    _agregar_tabla(doc, "Importancia SHAP (peso relativo %)", tablas.get("shap_importancia"))
    _agregar_tabla(doc, "Clústeres — tamaño (%)", tablas.get("cluster_tamanos"))
    _agregar_tabla(doc, "Clústeres — perfiles medios", tablas.get("cluster_perfiles"))

    grafico = resultado.get("grafico_shap")
    if grafico and Path(grafico).exists():
        doc.add_heading("Gráfico SHAP", level=2)
        doc.add_picture(str(grafico), width=Inches(6))

    doc.add_paragraph(
        "Nota metodológica: índice de exclusión digital compuesto (acceso hogar + uso individual TIC); "
        "movilidad social operada como proxies educativos, laborales y de ingreso. "
        "Fuente primaria: INDEC (www.indec.gob.ar)."
    )

    doc.save(path)
    return path


def exportar_word_bytes(
    resultado: dict,
    *,
    titulo: str,
    periodo: str,
    ambito: str,
) -> bytes:
    import tempfile

    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f:
        p = Path(f.name)
    try:
        exportar_word(resultado, p, titulo=titulo, periodo=periodo, ambito=ambito)
        return p.read_bytes()
    finally:
        p.unlink(missing_ok=True)
